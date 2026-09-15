"""El libro de Excel que se le manda a cada coordinador.

No toca la base de datos: entra el cuadro ya armado y sale el .xlsx.
"""
import os
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from services import reporte_posiciones  # noqa: E402


def voluntario(nombre, alimentacion=""):
    return {
        "voluntario": nombre,
        "telefono": "8091112222",
        "email": f"{nombre.split()[0].lower()}@example.com",
        "talla": "M",
        "tipo_sangre": "O+",
        "emergencia": "María Pérez (madre) 8093334444",
        "alimentacion": alimentacion,
    }


CUADRO = [
    {
        "puesto": "Hidratación y Snacks",
        "turnos": [
            {
                "dia": "2026-01-24",
                "turno": "A",
                "horario": "08:00-12:00",
                "cupos": 3,
                "equipo": [
                    voluntario("Ana Ortiz", "Refrigerio 09:30 · Almuerzo 12:00"),
                    voluntario("Luis Gómez", "Refrigerio 09:30"),
                ],
            },
            {
                "dia": "2026-01-24",
                "turno": "B",
                "horario": "12:00-16:00",
                "cupos": 1,
                "equipo": [voluntario("Ana Ortiz", "Refrigerio 14:30")],
            },
        ],
    },
    {
        "puesto": "Control de Vueltas",
        "turnos": [
            {
                "dia": "2026-01-24",
                "turno": "A",
                "horario": "08:00-12:00",
                "cupos": 2,
                "equipo": [voluntario("Sol Díaz", "Refrigerio 09:30")],
            },
        ],
    },
]


def abrir(cuadro=CUADRO, evento="Carrera Activa"):
    return load_workbook(reporte_posiciones.construir_libro(cuadro, evento))


class TestNombreDeHoja:
    def test_quita_lo_que_excel_no_admite(self):
        nombre = reporte_posiciones.nombre_de_hoja("Área de Carpas / Zona de Atletas", set())
        assert "/" not in nombre
        assert len(nombre) <= 31

    def test_dos_puestos_que_se_recortan_igual_no_chocan(self):
        usados = set()
        largo = "Puesto de trabajo con nombre larguísimo "
        primero = reporte_posiciones.nombre_de_hoja(largo + "uno", usados)
        segundo = reporte_posiciones.nombre_de_hoja(largo + "dos", usados)
        assert primero != segundo
        assert len(segundo) <= 31


class TestLibro:
    def test_una_hoja_por_puesto_mas_el_resumen(self):
        libro = abrir()
        assert libro.sheetnames == ["Resumen", "Hidratación y Snacks", "Control de Vueltas"]

    def test_el_resumen_cuenta_cupos_y_vacantes(self):
        hoja = abrir()["Resumen"]
        filas = {fila[0]: fila for fila in hoja.iter_rows(min_row=5, values_only=True)}
        # Hidratación: 4 cupos, 3 asignados, 1 vacante, 2 personas distintas
        assert filas["Hidratación y Snacks"][1:6] == (2, 4, 3, 1, 2)
        # Control de Vueltas: 2 cupos, 1 asignado, 1 vacante
        assert filas["Control de Vueltas"][1:6] == (1, 2, 1, 1, 1)

    def test_la_hoja_del_puesto_lleva_turno_equipo_y_comida(self):
        hoja = abrir()["Hidratación y Snacks"]
        texto = [fila[0] for fila in hoja.iter_rows(values_only=True)]
        assert texto[0] == "Hidratación y Snacks"
        assert any("Turno A" in (t or "") and "08:00-12:00" in (t or "") for t in texto)
        assert "Ana Ortiz" in texto and "Luis Gómez" in texto

        comidas = [fila[6] for fila in hoja.iter_rows(values_only=True)]
        assert "Refrigerio 09:30 · Almuerzo 12:00" in comidas

    def test_los_cupos_sin_cubrir_se_ven(self):
        hoja = abrir()["Hidratación y Snacks"]
        libres = [fila[0] for fila in hoja.iter_rows(values_only=True) if fila[0] == "(cupo libre)"]
        assert len(libres) == 1

    def test_un_puesto_sin_turnos_no_rompe_el_libro(self):
        libro = abrir([{"puesto": "Sin turnos", "turnos": []}])
        hoja = libro["Sin turnos"]
        assert hoja["A4"].value == "Este puesto no tiene turnos configurados"
