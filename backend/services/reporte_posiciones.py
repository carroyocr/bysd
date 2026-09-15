"""El cuadro de trabajo de cada puesto, en un libro de Excel.

Lo que se le manda al coordinador: su gente, a que hora entra cada uno, como
localizarlo y que comida le toca. Una hoja por puesto, para poder mandar la
suya a cada quien, y una de resumen al frente con los cupos que faltan.

Aqui no se habla con la base de datos: entra el cuadro ya armado y sale el
libro, para poder probarlo sin levantar nada.
"""
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CABECERAS_EQUIPO = [
    ("voluntario", "Voluntario", 30),
    ("telefono", "Teléfono", 16),
    ("email", "Correo", 32),
    ("talla", "Talla", 8),
    ("tipo_sangre", "Sangre", 9),
    ("emergencia", "Contacto de emergencia", 34),
    ("alimentacion", "Alimentación del turno", 34),
]

CABECERAS_RESUMEN = [
    ("puesto", "Puesto", 38),
    ("turnos", "Turnos", 10),
    ("cupos", "Cupos", 10),
    ("asignados", "Asignados", 12),
    ("vacantes", "Vacantes", 12),
    ("personas", "Personas", 12),
]

NARANJA = "E8772E"
GRIS = "F1F5F9"

_TITULO = Font(bold=True, size=14)
_CABECERA = Font(bold=True, color="FFFFFF")
_RELLENO_CABECERA = PatternFill("solid", fgColor=NARANJA)
_TURNO = Font(bold=True)
_RELLENO_TURNO = PatternFill("solid", fgColor=GRIS)
_VACANTE = Font(italic=True, color="B91C1C")
_BORDE_FINO = Border(bottom=Side(style="thin", color="D1D5DB"))

# Excel no admite estos caracteres en el nombre de una hoja, y la corta a 31
PROHIBIDOS_EN_HOJA = re.compile(r"[\\/*?:\[\]]")


def nombre_de_hoja(titulo: str, usados: set) -> str:
    """Nombre valido y unico para la hoja de un puesto.

    Los puestos llevan barras y son largos ("Área de Carpas / Zona de
    Atletas"): Excel no lo acepta tal cual y al recortarlos a 31 caracteres dos
    puestos distintos pueden quedar iguales, asi que el repetido lleva numero.
    """
    limpio = PROHIBIDOS_EN_HOJA.sub("-", titulo or "").strip() or "Puesto"
    nombre = limpio[:31]
    if nombre.lower() not in usados:
        usados.add(nombre.lower())
        return nombre

    for i in range(2, 100):
        sufijo = f" ({i})"
        candidato = limpio[: 31 - len(sufijo)] + sufijo
        if candidato.lower() not in usados:
            usados.add(candidato.lower())
            return candidato
    return nombre


def _escribir_cabecera(hoja, fila: int, cabeceras) -> int:
    for columna, (_, titulo, _ancho) in enumerate(cabeceras, start=1):
        celda = hoja.cell(row=fila, column=columna, value=titulo)
        celda.font = _CABECERA
        celda.fill = _RELLENO_CABECERA
        celda.alignment = Alignment(vertical="center")
    return fila + 1


def _anchos(hoja, cabeceras):
    for columna, (_, _titulo, ancho) in enumerate(cabeceras, start=1):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho


def _hoja_resumen(libro, posiciones, evento: str):
    hoja = libro.active
    hoja.title = "Resumen"

    hoja["A1"] = "Voluntarios por posición"
    hoja["A1"].font = _TITULO
    hoja["A2"] = evento
    hoja["A2"].font = Font(color="6B7280")

    fila = _escribir_cabecera(hoja, 4, CABECERAS_RESUMEN)
    for posicion in posiciones:
        cupos = sum(t["cupos"] for t in posicion["turnos"])
        asignados = sum(len(t["equipo"]) for t in posicion["turnos"])
        personas = {v["email"] for t in posicion["turnos"] for v in t["equipo"] if v.get("email")}
        valores = {
            "puesto": posicion["puesto"],
            "turnos": len(posicion["turnos"]),
            "cupos": cupos,
            "asignados": asignados,
            "vacantes": max(cupos - asignados, 0),
            "personas": len(personas),
        }
        for columna, (clave, _titulo, _ancho) in enumerate(CABECERAS_RESUMEN, start=1):
            celda = hoja.cell(row=fila, column=columna, value=valores[clave])
            celda.border = _BORDE_FINO
            if clave == "vacantes" and valores[clave]:
                celda.font = _VACANTE
        fila += 1

    _anchos(hoja, CABECERAS_RESUMEN)
    hoja.freeze_panes = "A5"


def _hoja_posicion(libro, posicion, evento: str, usados: set):
    hoja = libro.create_sheet(nombre_de_hoja(posicion["puesto"], usados))

    hoja["A1"] = posicion["puesto"]
    hoja["A1"].font = _TITULO
    hoja["A2"] = evento
    hoja["A2"].font = Font(color="6B7280")

    fila = 4
    for turno in posicion["turnos"]:
        # La franja, de una pieza, para que se lea de un vistazo antes de la
        # lista de quien la cubre.
        libres = max(turno["cupos"] - len(turno["equipo"]), 0)
        encabezado = (
            f"{turno['dia']} · Turno {turno['turno']} · {turno['horario']}"
            f"  ({len(turno['equipo'])} de {turno['cupos']} cupos)"
        )
        celda = hoja.cell(row=fila, column=1, value=encabezado)
        celda.font = _TURNO
        for columna in range(1, len(CABECERAS_EQUIPO) + 1):
            hoja.cell(row=fila, column=columna).fill = _RELLENO_TURNO
        fila += 1

        fila = _escribir_cabecera(hoja, fila, CABECERAS_EQUIPO)

        for voluntario in turno["equipo"]:
            for columna, (clave, _titulo, _ancho) in enumerate(CABECERAS_EQUIPO, start=1):
                celda = hoja.cell(row=fila, column=columna, value=voluntario.get(clave, ""))
                celda.border = _BORDE_FINO
                celda.alignment = Alignment(wrap_text=(clave in ("emergencia", "alimentacion")), vertical="top")
            fila += 1

        for _ in range(libres):
            celda = hoja.cell(row=fila, column=1, value="(cupo libre)")
            celda.font = _VACANTE
            celda.border = _BORDE_FINO
            fila += 1

        fila += 1  # una linea en blanco entre turnos

    if not posicion["turnos"]:
        hoja.cell(row=4, column=1, value="Este puesto no tiene turnos configurados")

    _anchos(hoja, CABECERAS_EQUIPO)


def construir_libro(posiciones: list, evento: str = "") -> io.BytesIO:
    """Devuelve el .xlsx listo para descargar."""
    libro = Workbook()
    _hoja_resumen(libro, posiciones, evento)

    usados = {"resumen"}
    for posicion in posiciones:
        _hoja_posicion(libro, posicion, evento, usados)

    memoria = io.BytesIO()
    libro.save(memoria)
    memoria.seek(0)
    return memoria
